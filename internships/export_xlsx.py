#!/usr/bin/env python3
"""
internships/export_xlsx.py
--------------------------
Rebuilds the Internship_Tracker workbook from the Recruiting Desk data, in the
same layout as the original hand-built tracker: a Dashboard tab with live
formulas, a Tracker tab with the 18 original columns, dropdowns, Days Left /
Days Since Applied formulas and the red / orange / yellow warning rules, plus a
New Leads tab for postings the bot found that you have not acted on.

    python3 internships/export_xlsx.py DOCS_DIR OUT.xlsx

DOCS_DIR is where `ArtifactData list ... out_dir=DOCS_DIR` saved the tracker:
DOCS_DIR/apps/*.json and DOCS_DIR/leads/*.json.

Needs openpyxl (`pip install openpyxl`).
"""
import datetime as dt
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# Vocabulary shared with desk.js and tracker.html.
STAGES = ["Not Open Yet", "Researching", "Interested", "Applying", "Applied", "Online Assessment",
          "Interviewing", "Final Round", "Offer", "Rejected", "Withdrawn", "Closed"]
WATCH = {"Not Open Yet", "Researching", "Interested", "Applying"}
DONE = {"Rejected", "Withdrawn", "Closed"}
CATEGORIES = ["Asset Management", "Wealth Mgmt / Planning", "Corporate Finance", "Operations", "Banking",
              "Markets", "Risk", "Other"]
SOURCES = ["12twenty", "LinkedIn", "Company Site", "Career Fair", "Handshake", "Referral", "Simplify", "Other"]
TIERS = ["Lead", "Target", "Reach"]

# Colors lifted from the original workbook.
NAVY = "FF1F3864"
GRAY = "FFF2F2F2"
STAGE_FILL = {
    "Not Open Yet": "FFEDEDED", "Researching": "FFDDEBF7", "Interested": "FFBDD7EE", "Applying": "FFFFF2CC",
    "Applied": "FF9BC2E6", "Online Assessment": "FFE4DFEC", "Interviewing": "FFF8CBAD", "Final Round": "FFF4B084",
    "Offer": "FFC6EFCE", "Rejected": "FFFFC7CE", "Withdrawn": "FFD9D9D9", "Closed": "FFD9D9D9",
}
HEADERS = ["#", "Company", "Role", "Location", "Category", "Fit", "Source", "Status", "Deadline", "Days Left",
           "Date Applied", "Days Since Applied", "Next Step", "Follow-Up Date", "Contact", "Pay", "Link", "Notes"]
WIDTHS = [5, 22, 40, 24, 20, 9, 13, 17, 13, 10, 13, 11, 44, 13, 26, 22, 16, 60]
FIRST, LAST = 5, 200          # data rows, as in the original
DATE_FMT = "mmm d, yyyy"


def load(docs_dir, collection):
    rows = []
    for f in sorted(glob.glob(os.path.join(docs_dir, collection, "*.json"))):
        with open(f, encoding="utf-8") as fh:
            d = json.load(fh)
        d.setdefault("key", os.path.basename(f)[:-5])
        rows.append(d)
    return rows


def as_date(v):
    try:
        return dt.date.fromisoformat(str(v)[:10]) if v else None
    except ValueError:
        return None


def order(a):
    s = a.get("stage", "Applied")
    group = 0 if s == "Offer" else 1 if s not in WATCH and s not in DONE else 2 if s in WATCH else 3
    rank = STAGES.index(s) if s in STAGES else 4
    return (group, -rank, str(a.get("company", "")).lower())


def header_row(ws, row, labels):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def build_tracker(wb, apps):
    ws = wb.create_sheet("Tracker")
    ws["A1"] = "Internship Application Tracker — Summer 2027"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = ("Exported from the Recruiting Desk. The live copy is the Recruiting Desk page; run /internships in "
                "Claude Code to refresh it. Gray columns calculate automatically. Re-check postings before applying.")
    ws["A2"].font = Font(italic=True, color="FF595959")
    header_row(ws, 4, HEADERS)
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w

    for n, a in enumerate(sorted(apps, key=order), start=1):
        r = FIRST + n - 1
        deadline = as_date(a.get("deadline"))
        values = [n, a.get("company", ""), a.get("role", ""), a.get("location", ""), a.get("category", ""),
                  a.get("tier", ""), a.get("source", ""), a.get("stage", ""),
                  deadline or a.get("deadline_note", "") or None, f'=IF(ISNUMBER(I{r}),I{r}-TODAY(),"")',
                  as_date(a.get("applied")), f'=IF(ISNUMBER(K{r}),TODAY()-K{r},"")', a.get("next", ""),
                  as_date(a.get("followup")), a.get("contact", ""), a.get("pay", ""), None, a.get("notes", "")]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=v if v != "" else None)
            c.alignment = Alignment(vertical="top", wrap_text=col in (3, 13, 18))
            if isinstance(v, dt.date):
                c.number_format = DATE_FMT
        link = a.get("link") or ""
        if link.startswith("http"):
            c = ws.cell(row=r, column=17, value="Open posting")
            c.hyperlink = link
            c.font = Font(color="FF0563C1", underline="single")
    for r in range(FIRST, LAST + 1):
        for col in (10, 12):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=GRAY)

    for col, items in (("E", CATEGORIES), ("F", TIERS), ("G", SOURCES), ("H", STAGES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
        dv.add(f"{col}{FIRST}:{col}{LAST}")
        ws.add_data_validation(dv)

    rng = f"H{FIRST}:H{LAST}"
    for s, color in STAGE_FILL.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$H{FIRST}="{s}"'], fill=PatternFill("solid", bgColor=color)))
    watch = f'OR($H{FIRST}="Not Open Yet",$H{FIRST}="Researching",$H{FIRST}="Interested",$H{FIRST}="Applying")'
    ws.conditional_formatting.add(f"I{FIRST}:J{LAST}", FormulaRule(
        formula=[f"AND(ISNUMBER($I{FIRST}),$I{FIRST}-TODAY()<=7,{watch})"],
        fill=PatternFill("solid", bgColor="FFFFC7CE"), font=Font(bold=True, color="FF9C0006")))
    ws.conditional_formatting.add(f"L{FIRST}:L{LAST}", FormulaRule(
        formula=[f'AND(ISNUMBER($L{FIRST}),$L{FIRST}>=14,$H{FIRST}="Applied")'],
        fill=PatternFill("solid", bgColor="FFFFEB9C"), font=Font(bold=True, color="FF9C5700")))
    ws.conditional_formatting.add(f"N{FIRST}:N{LAST}", FormulaRule(
        formula=[f'AND(ISNUMBER($N{FIRST}),$N{FIRST}<=TODAY(),$H{FIRST}<>"Offer",$H{FIRST}<>"Rejected",'
                 f'$H{FIRST}<>"Withdrawn",$H{FIRST}<>"Closed")'],
        fill=PatternFill("solid", bgColor="FFFCE4D6"), font=Font(bold=True, color="FFC65911")))
    ws.freeze_panes = "C5"
    return ws


def build_dashboard(wb):
    ws = wb.active
    ws.title = "Dashboard"
    for col, w in zip("ABCDE", (3, 34, 14, 3, 95)):
        ws.column_dimensions[col].width = w
    ws["B2"] = "Internship Search Dashboard"
    ws["B2"].font = Font(size=18, bold=True, color=NAVY)
    ws["B3"] = '="As of "&TEXT(TODAY(),"mmmm d, yyyy")'
    ws["B3"].font = Font(italic=True, color="FF595959")
    for cell, text in (("B5", "Key Numbers"), ("E5", "How to use this tracker"), ("B16", "Pipeline by Status")):
        ws[cell] = text
        ws[cell].font = Font(bold=True, size=12, color=NAVY)

    t = "Tracker!$H$5:$H$200"
    key_rows = [
        ("Companies tracked", "=COUNTA(Tracker!$B$5:$B$200)", "0"),
        ("Applications submitted", "=SUM(C21:C27)", "0"),
        ("Interviews (incl. OA & final round)", "=C22+C23+C24", "0"),
        ("Offers", "=C25", "0"),
        ("Response rate (heard back / applied)", '=IF(C7=0,"-",(C22+C23+C24+C25+C26)/C7)', "0%"),
        ("Deadlines in next 7 days (not applied)",
         "=SUMPRODUCT(ISNUMBER(Tracker!$I$5:$I$200)*(Tracker!$I$5:$I$200>=TODAY())*(Tracker!$I$5:$I$200<=TODAY()+7)*"
         f'(({t}="Not Open Yet")+({t}="Researching")+({t}="Interested")+({t}="Applying")))', "0"),
        ("Follow-ups due today or overdue",
         "=SUMPRODUCT(ISNUMBER(Tracker!$N$5:$N$200)*(Tracker!$N$5:$N$200<=TODAY())*"
         f'({t}<>"Offer")*({t}<>"Rejected")*({t}<>"Withdrawn")*({t}<>"Closed")*(Tracker!$B$5:$B$200<>""))', "0"),
        ("Not open yet — watch list", f'=COUNTIF({t},"Not Open Yet")', "0"),
    ]
    for i, (label, formula, fmt) in enumerate(key_rows):
        ws.cell(row=6 + i, column=2, value=label)
        c = ws.cell(row=6 + i, column=3, value=formula)
        c.number_format = fmt
        c.font = Font(bold=True)
    for cell in ("C11", "C12"):
        ws.conditional_formatting.add(cell, FormulaRule(formula=[f"{cell}>0"], fill=PatternFill("solid", bgColor="FFFFC7CE"),
                                                        font=Font(bold=True, color="FF9C0006")))
    for i, s in enumerate(STAGES):
        ws.cell(row=17 + i, column=2, value=s).fill = PatternFill("solid", fgColor=STAGE_FILL[s])
        ws.cell(row=17 + i, column=3, value=f"=COUNTIF({t},B{17 + i})")

    how = [
        "1. The live tracker is the Recruiting Desk page on claude.ai. This file is a snapshot of it.",
        "2. Type /internships in Claude Code (phone, Mac or desktop) to pull new 12twenty, LinkedIn and SimplifyJobs",
        "   postings and log new applications from your confirmation emails. Nothing runs until you type it.",
        "3. Change Status on the page (or here) every time something moves. The bot keeps your edits.",
        "4. The day you apply, Follow-Up Date is set 12 days out. Orange = send a short check-in.",
        "5. Red deadline = due within a week and not applied yet. Yellow Days Since Applied = 14+ days, no reply.",
        "",
        "Status meanings:",
        "Not Open Yet = company recruits, 2027 posting not live · Researching = checking fit/openings",
        "Interested = open, plan to apply · Applying = working on it now · Applied = submitted",
        "Online Assessment = HireVue / test · Interviewing / Final Round · Offer · Rejected",
        "Withdrawn = you pulled out · Closed = missed or posting removed",
    ]
    for i, line in enumerate(how):
        ws.cell(row=6 + i, column=5, value=line)
    return ws


def build_leads(wb, leads, apps):
    ws = wb.create_sheet("New Leads")
    applied = {a.get("key") for a in apps}
    rows = [l for l in leads if l.get("key") not in applied and l.get("status", "new") in ("new", "saved")]
    fit_rank = {"strong": 0, "possible": 1, "low": 2}
    rows.sort(key=lambda l: (fit_rank.get(l.get("fit"), 3), l.get("deadline") or "9999"))
    ws["A1"] = "New Leads — found by /internships, not yet in your tracker"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    hdr = ["Company", "Role", "Fit", "Why it matched", "Deadline", "Days Left", "Location", "Source", "Found", "Link"]
    header_row(ws, 3, hdr)
    for col, w in zip("ABCDEFGHIJ", (24, 48, 10, 36, 13, 10, 24, 11, 12, 14)):
        ws.column_dimensions[col].width = w
    for i, l in enumerate(rows):
        r = 4 + i
        vals = [l.get("company"), l.get("role"), l.get("fit"), l.get("why"), as_date(l.get("deadline")),
                f'=IF(ISNUMBER(E{r}),E{r}-TODAY(),"")', l.get("location"), l.get("source"), as_date(l.get("found")), None]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v or None)
            if isinstance(v, dt.date):
                c.number_format = DATE_FMT
        if str(l.get("link", "")).startswith("http"):
            c = ws.cell(row=r, column=10, value="Open")
            c.hyperlink = l["link"]
            c.font = Font(color="FF0563C1", underline="single")
    ws.freeze_panes = "B4"
    return len(rows)


def main(argv):
    if len(argv) != 3:
        sys.exit(__doc__)
    docs_dir, out = argv[1], argv[2]
    apps, leads = load(docs_dir, "apps"), load(docs_dir, "leads")
    wb = Workbook()
    build_dashboard(wb)
    build_tracker(wb, apps)
    n_leads = build_leads(wb, leads, apps)
    wb.save(out)
    print(f"Wrote {out}: {len(apps)} tracker rows, {n_leads} open leads.")


if __name__ == "__main__":
    main(sys.argv)
